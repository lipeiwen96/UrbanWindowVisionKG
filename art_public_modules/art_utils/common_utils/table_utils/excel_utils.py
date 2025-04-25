"""
Excel表格处理工具
@建筑研究 李沛文
"""
from dataclasses import field, dataclass
from openpyxl.utils import get_column_letter, column_index_from_string


@dataclass(order=True, unsafe_hash=True)
class ExcelMethods:
    @staticmethod
    def get_column_letter_by_index(column_index: int):
        return get_column_letter(column_index)

    @staticmethod
    def get_column_index_by_string(column_string: str):
        return column_index_from_string(column_string)

    @staticmethod
    def merge_and_write(worksheet, start_cell: str, end_cell: str, content, format):
        """
        xlsxwriter库中merge_range的进阶版，可以合并单个单元格并添加内容
        """
        if start_cell == end_cell:
            worksheet.write(start_cell, content, format)
        else:
            worksheet.merge_range(f"{start_cell}:{end_cell}", content, format)


if __name__ == "__main__":
    print("1:", ExcelMethods.get_column_letter_by_index(1))
    print("25:", ExcelMethods.get_column_letter_by_index(25))
    print("100:", ExcelMethods.get_column_letter_by_index(100))
    print("B:", ExcelMethods.get_column_index_by_string("B"))
    print("Cc:", ExcelMethods.get_column_index_by_string("Cc"))
    print("hb:", ExcelMethods.get_column_index_by_string("hb"))
